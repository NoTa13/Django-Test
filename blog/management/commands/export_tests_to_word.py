import os
from django.core.management.base import BaseCommand
from docx import Document

class Command(BaseCommand):
    help = 'Экспортирует код тестов из файла tests.py в файл Word (.docx)'

    def handle(self, *args, **options):
        # Путь к файлу с тестами
        tests_path = os.path.join('blog', 'tests.py')
        output_filename = 'Django_Tests_Report.docx'

        if not os.path.exists(tests_path):
            self.stdout.write(self.style.ERROR(f'Файл {tests_path} не найден!'))
            return

        # Читаем содержимое tests.py
        with open(tests_path, 'r', encoding='utf-8') as f:
            code_content = f.read()

        # Создаем Word документ
        doc = Document()
        doc.add_heading('Отчет о юнит-тестах Django проекта', level=0)
        
        doc.add_heading('Содержимое файла blog/tests.py', level=1)
        doc.add_paragraph('Ниже представлен исходный код всех разработанных модульных тестов:')

        # Добавляем код с моноширинным шрифтом
        p = doc.add_paragraph()
        run = p.add_run(code_content)
        run.font.name = 'Courier New'

        # Сохраняем файл
        doc.save(output_filename)
        self.stdout.write(self.style.SUCCESS(f'Успешно создан Word-файл: {output_filename}'))